"""Result-grid export engine — CSV / XLSX / XML / Markdown / JSON / HTML.

Pure functions, no Qt: the UI collects (columns, rows) from the grid and
calls :func:`export_rows`. All writers take ``rows`` as a list of lists of
cell text (already formatted by the grid layer).
"""
from __future__ import annotations

import csv
import io
import json
from typing import Any, List, Optional, Sequence, Tuple
from xml.sax.saxutils import escape as xml_escape

#: (key, label-en, extension) — label is localized by the UI layer.
FORMATS: Tuple[Tuple[str, str, str], ...] = (
    ("csv", "CSV", "csv"),
    ("xlsx", "Excel (XLSX)", "xlsx"),
    ("xml", "XML", "xml"),
    ("markdown", "Markdown", "md"),
    ("json", "JSON", "json"),
    ("html", "HTML", "html"),
)

DEFAULT_DELIMITER = ","


def extension(fmt: str) -> str:
    for key, _label, ext in FORMATS:
        if key == fmt:
            return ext
    raise ValueError(f"unknown export format: {fmt}")


def _clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def _clean_rows(rows: Sequence[Sequence[Any]]) -> List[List[str]]:
    return [[_clean(v) for v in row] for row in rows]


# ---------------------------------------------------------------------------
# per-format writers
# ---------------------------------------------------------------------------

def to_csv(columns: Sequence[str], rows: Sequence[Sequence[Any]],
           delimiter: str = DEFAULT_DELIMITER) -> str:
    buf = io.StringIO()
    # utf-8-sig equivalent handled by caller at file-write time.
    writer = csv.writer(buf, delimiter=delimiter, lineterminator="\n")
    writer.writerow(list(columns))
    writer.writerows(_clean_rows(rows))
    return buf.getvalue()


def to_xlsx(path: str, columns: Sequence[str],
            rows: Sequence[Sequence[Any]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "results"
    ws.append(list(columns))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in _clean_rows(rows):
        ws.append(row)
    # rough auto-width, capped
    for i, col in enumerate(columns, start=1):
        width = max(len(str(col)) + 2, 10)
        for row in rows:
            idx = i - 1
            if idx < len(row) and row[idx] is not None:
                width = max(width, min(len(str(row[idx])) + 2, 60))
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    wb.save(path)


def to_xml(columns: Sequence[str], rows: Sequence[Sequence[Any]]) -> str:
    parts = ['<?xml version="1.0" encoding="UTF-8"?>',
             '<result>']
    clean = _clean_rows(rows)
    for row in clean:
        parts.append("  <row>")
        for i, col in enumerate(columns):
            value = row[i] if i < len(row) else ""
            parts.append(f'    <field name="{xml_escape(str(col))}">'
                         f'{xml_escape(value)}</field>')
        parts.append("  </row>")
    parts.append("</result>")
    return "\n".join(parts) + "\n"


def to_markdown(columns: Sequence[str], rows: Sequence[Sequence[Any]]) -> str:
    def cell(text: str) -> str:
        return text.replace("|", "\\|").replace("\n", " ")

    lines = ["| " + " | ".join(cell(str(c)) for c in columns) + " |",
             "|" + "|".join(" --- " for _ in columns) + "|"]
    for row in _clean_rows(rows):
        lines.append("| " + " | ".join(cell(v) for v in row) + " |")
    return "\n".join(lines) + "\n"


def to_json(columns: Sequence[str], rows: Sequence[Sequence[Any]]) -> str:
    array = []
    for row in rows:
        d = {}
        for i, col in enumerate(columns):
            d[str(col)] = row[i] if i < len(row) else None
        array.append(d)
    return json.dumps(array, ensure_ascii=False, indent=2, default=str) + "\n"


def to_html(columns: Sequence[str], rows: Sequence[Sequence[Any]]) -> str:
    parts = [
        "<!DOCTYPE html>",
        '<html><head><meta charset="utf-8">',
        "<title>InfluxDB Manager export</title>",
        "<style>",
        "body{font-family:system-ui,sans-serif;margin:24px;}",
        "table{border-collapse:collapse;}",
        "th,td{border:1px solid #ccc;padding:4px 10px;text-align:left;}",
        "th{background:#f0f4f8;}",
        "td.num{text-align:right;font-variant-numeric:tabular-nums;}",
        "</style></head><body>",
        "<table>",
        "<thead><tr>" + "".join(
            f"<th>{xml_escape(str(c))}</th>" for c in columns) + "</tr></thead>",
        "<tbody>",
    ]
    for row in _clean_rows(rows):
        tds = []
        for i in range(len(columns)):
            value = row[i] if i < len(row) else ""
            cls = ' class="num"' if _looks_numeric(value) else ""
            tds.append(f"<td{cls}>{xml_escape(value)}</td>")
        parts.append("<tr>" + "".join(tds) + "</tr>")
    parts.append("</tbody></table></body></html>")
    return "\n".join(parts) + "\n"


def _looks_numeric(text: str) -> bool:
    if not text:
        return False
    try:
        float(text)
        return True
    except ValueError:
        return False


# ---------------------------------------------------------------------------
# entry point
# ---------------------------------------------------------------------------

def export_rows(path: str, fmt: str, columns: Sequence[str],
                rows: Sequence[Sequence[Any]],
                delimiter: str = DEFAULT_DELIMITER) -> None:
    """Write ``rows`` to ``path`` in the given format.

    XLSX is written directly by openpyxl; every other format produces text
    (CSV is written with utf-8-sig so Excel opens Chinese headers correctly).
    """
    if fmt == "csv":
        if len(delimiter) != 1:
            raise ValueError(f"CSV delimiter must be a single character, "
                             f"got {delimiter!r}")
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            f.write(to_csv(columns, rows, delimiter))
    elif fmt == "xlsx":
        to_xlsx(path, columns, rows)
    elif fmt == "xml":
        with open(path, "w", encoding="utf-8") as f:
            f.write(to_xml(columns, rows))
    elif fmt == "markdown":
        with open(path, "w", encoding="utf-8") as f:
            f.write(to_markdown(columns, rows))
    elif fmt == "json":
        with open(path, "w", encoding="utf-8") as f:
            f.write(to_json(columns, rows))
    elif fmt == "html":
        with open(path, "w", encoding="utf-8") as f:
            f.write(to_html(columns, rows))
    else:
        raise ValueError(f"unknown export format: {fmt}")
