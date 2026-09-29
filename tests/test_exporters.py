"""Exporter engine unit tests — pure functions, no Qt, no network."""
from __future__ import annotations

import csv
import io
import json
import os
import xml.etree.ElementTree as ET

import pytest

from net.sakurain.influxdbstudio.core import exporters

COLUMNS = ["time", "name", "value", "note"]
ROWS = [
    ["2026-09-27T16:00:00Z", "alpha", "1.5", "含中文,逗号"],
    ["2026-09-27T16:01:00Z", "beta", None, 'say "hi"'],
    [None, "gamma|delta", "3", ""],
]

CLEAN = [["2026-09-27T16:00:00Z", "alpha", "1.5", "含中文,逗号"],
         ["2026-09-27T16:01:00Z", "beta", "", 'say "hi"'],
         ["", "gamma|delta", "3", ""]]


class TestFormats:
    def test_csv_default_comma(self):
        text = exporters.to_csv(COLUMNS, ROWS)
        parsed = list(csv.reader(io.StringIO(text)))
        assert parsed[0] == COLUMNS
        assert parsed[1] == CLEAN[0]
        assert len(parsed) == 4

    def test_csv_custom_delimiter(self):
        text = exporters.to_csv(COLUMNS, ROWS, delimiter=";")
        parsed = list(csv.reader(io.StringIO(text), delimiter=";"))
        assert parsed[1] == CLEAN[0]

    def test_csv_quoting_specials(self):
        text = exporters.to_csv(COLUMNS, ROWS)
        # comma inside a field must be quoted
        assert '"含中文,逗号"' in text
        assert '"say ""hi"""' in text

    def test_csv_bad_delimiter_raises(self, tmp_path):
        with pytest.raises(ValueError):
            exporters.export_rows(str(tmp_path / "x.csv"), "csv",
                                  COLUMNS, ROWS, delimiter="::")

    def test_export_rows_csv_uses_bom(self, tmp_path):
        p = str(tmp_path / "out.csv")
        exporters.export_rows(p, "csv", COLUMNS, ROWS)
        raw = open(p, "rb").read()
        assert raw.startswith(b"\xef\xbb\xbf")

    def test_xlsx_roundtrip(self, tmp_path):
        p = str(tmp_path / "out.xlsx")
        exporters.export_rows(p, "xlsx", COLUMNS, ROWS)
        from openpyxl import load_workbook
        wb = load_workbook(p)
        ws = wb.active
        grid = [[("" if c.value is None else str(c.value)) for c in row]
                for row in ws.iter_rows()]
        assert grid[0] == COLUMNS
        assert grid[1] == CLEAN[0]
        assert len(grid) == 4

    def test_xml_wellformed_and_escaped(self, tmp_path):
        p = str(tmp_path / "out.xml")
        exporters.export_rows(p, "xml", COLUMNS, ROWS)
        root = ET.parse(p).getroot()
        assert root.tag == "result"
        rows = root.findall("row")
        assert len(rows) == 3
        fields = rows[0].findall("field")
        assert [f.get("name") for f in fields] == COLUMNS
        assert fields[3].text == "含中文,逗号"
        # pipe in markdown-special value must not break xml; & < > escaped
        text = open(p, encoding="utf-8").read()
        assert "gamma|delta" in text

    def test_markdown_table(self):
        text = exporters.to_markdown(COLUMNS, ROWS)
        lines = text.strip().splitlines()
        assert lines[0].startswith("| time |")
        assert lines[1].startswith("|") and "---" in lines[1]
        assert len(lines) == 5  # header + separator + 3 rows
        # pipe inside value escaped
        assert "gamma\\|delta" in text

    def test_json_typed_and_ascii(self):
        text = exporters.to_json(COLUMNS, ROWS)
        data = json.loads(text)
        assert len(data) == 3
        assert data[0]["note"] == "含中文,逗号"
        assert data[1]["value"] is None
        assert "含中文" in text  # ensure_ascii=False

    def test_html_escapes_and_numeric_class(self, tmp_path):
        p = str(tmp_path / "out.html")
        exporters.export_rows(p, "html", COLUMNS, ROWS)
        text = open(p, encoding="utf-8").read()
        assert "<table>" in text and "</html>" in text
        assert 'class="num"' in text  # "1.5" cell right-aligned
        assert "gamma|delta" in text

    def test_all_six_formats_accepted(self, tmp_path):
        for key, _label, ext in exporters.FORMATS:
            p = str(tmp_path / f"out.{ext}")
            exporters.export_rows(p, key, COLUMNS, ROWS)
            assert os.path.getsize(p) > 0, key

    def test_unknown_format_raises(self, tmp_path):
        with pytest.raises(ValueError):
            exporters.export_rows(str(tmp_path / "x.xyz"), "xyz", COLUMNS, ROWS)

    def test_extension_lookup(self):
        assert exporters.extension("markdown") == "md"
        with pytest.raises(ValueError):
            exporters.extension("nope")
